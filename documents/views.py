from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.template.loader import render_to_string
from django.utils import timezone
from xhtml2pdf import pisa
from docx import Document as DocxDocument
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .models import DocumentModele, DemandeConge
from .forms import CertificatScolariteForm, DemandeCongeForm

# ---------------------------------------------------------------------------
# Catalogue des documents : un exemplaire PDF est disponible pour chaque type.
# ---------------------------------------------------------------------------
DOCUMENTS_CATALOG = [
    {
        'key': 'certificat_scolarite',
        'title': 'Certificat de scolarité',
        'description': "Atteste qu'un élève est régulièrement inscrit dans l'établissement.",
        'icon': 'fa-file-signature', 'color': '#059669',
        'generate_url': '/documents/certificat-scolarite/',
        'category': 'Administratif',
    },
    {
        'key': 'certificat_frequentation',
        'title': 'Certificat de fréquentation',
        'description': "Confirme la présence régulière de l'élève durant l'année.",
        'icon': 'fa-user-check', 'color': '#2563eb',
        'generate_url': '/documents/certificat-frequentation/',
        'category': 'Administratif',
    },
    {
        'key': 'certificat_transfert',
        'title': "Certificat de transfert",
        'description': "Délivré lorsqu'un élève quitte l'établissement pour une autre école.",
        'icon': 'fa-right-left', 'color': '#7c3aed',
        'generate_url': '/documents/certificat-transfert/',
        'category': 'Administratif',
    },
    {
        'key': 'attestation_reussite',
        'title': "Attestation de réussite",
        'description': "Officialise la réussite de l'élève à l'issue de l'année académique.",
        'icon': 'fa-award', 'color': '#d97706',
        'sample_only': True,
        'category': 'Administratif',
    },
    {
        'key': 'tableau_honneur',
        'title': 'Tableau d’honneur',
        'description': "Récompense les élèves ayant obtenu d'excellents résultats.",
        'icon': 'fa-medal', 'color': '#dc2626',
        'sample_only': True,
        'category': 'Pédagogique',
    },
    {
        'key': 'registre_presence',
        'title': 'Registre de présence',
        'description': 'Feuille d’émargement des présences pour une classe.',
        'icon': 'fa-clipboard-list', 'color': '#0f766e',
        'generate_url': '/documents/registre-presence/',
        'category': 'Pédagogique',
    },
    {
        'key': 'convocation_parent',
        'title': 'Convocation des parents',
        'description': 'Convoque officiellement un parent pour un entretien.',
        'icon': 'fa-envelope-open-text', 'color': '#b45309',
        'sample_only': True,
        'category': 'Administratif',
    },
    {
        'key': 'recu_paiement',
        'title': 'Reçu de paiement',
        'description': 'Justificatif numéroté des frais réglés par le parent.',
        'icon': 'fa-receipt', 'color': '#059669',
        'sample_only': True,
        'category': 'Financier',
    },
    {
        'key': 'bulletin',
        'title': 'Bulletin de notes',
        'description': 'Moyennes, coefficients, rang et appréciations de la période.',
        'icon': 'fa-file-lines', 'color': '#2563eb',
        'generate_url': '/grades/bulletins/',
        'category': 'Pédagogique',
    },
    {
        'key': 'demande_conge',
        'title': "Demande de congé (personnel)",
        'description': 'Formulaire de demande de congé pour le personnel.',
        'icon': 'fa-plane-departure', 'color': '#7c3aed',
        'generate_url': '/documents/demande-conge/',
        'category': 'RH',
    },
]


# ---------------------------------------------------------------------------
# Données d'exemple fictives par type de document (PDF et Word).
# ---------------------------------------------------------------------------
DOC_SAMPLE_DATA = {
        'certificat_scolarite': {
            'intro': ("Le Directeur de l'établissement atteste que l'élève dont "
                      "l'identité suit est régulièrement inscrit pour l'année :"),
            'fields': [
                ('Nom & Prénom', 'ADJOVI Koffi Marc'),
                ('Matricule', 'STD-2026-0125'),
                ('Né(e) le', '12/03/2012 à Cotonou'),
                ('Classe', '6ème A'),
                ('Année scolaire', '2025-2026'),
            ],
            'conclusion': "En foi de quoi, le présent certificat est délivré à l'intéressé pour servir et valoir ce que de droit.",
        },
        'certificat_frequentation': {
            'intro': "Le Directeur atteste que l'élève ci-dessous fréquente régulièrement l'établissement :",
            'fields': [
                ('Nom & Prénom', 'TRAORE Aminata'),
                ('Matricule', 'STD-2026-0087'),
                ('Classe', '5ème B'),
                ('Taux de fréquentation', '94 %'),
                ('Année scolaire', '2025-2026'),
            ],
            'conclusion': "Le présent certificat est délivré pour faire valoir ce que de droit.",
        },
        'certificat_transfert': {
            'intro': "Le Directeur atteste que l'élève ci-dessous a quitté l'établissement pour transfert :",
            'fields': [
                ('Nom & Prénom', 'KEITA Ibrahim'),
                ('Matricule', 'STD-2025-0042'),
                ('Dernière classe fréquentée', '4ème A'),
                ('École d’accueil', 'Groupe Scolaire La Réussite'),
                ('Motif', 'Déménagement de la famille'),
            ],
            'conclusion': "Son dossier scolaire sera transmis sur simple demande de l'établissement d'accueil.",
        },
        'attestation_reussite': {
            'intro': "Le Directeur atteste que l'élève ci-dessous a réussi sa classe avec honneur :",
            'fields': [
                ('Nom & Prénom', 'SANOGO Rokia'),
                ('Matricule', 'STD-2026-0033'),
                ('Classe validée', 'Terminale C'),
                ('Moyenne annuelle', '14,75 / 20'),
                ('Décision du conseil', 'Admise'),
            ],
            'conclusion': "Félicitations du Conseil de Discipline pour la conduite exemplaire de l'élève.",
        },
        'tableau_honneur': {
            'intro': "Le tableau d'honneur du trimestre récompense les élèves suivants :",
            'fields': [
                ('Élève', 'COULIBALY Mariam — 1ère A'),
                ('Moyenne du trimestre', '17,25 / 20'),
                ('Rang', '1er / 32'),
                ('Mention', 'Très Bien'),
                ('Trimestre', '1er Trimestre — 2025-2026'),
            ],
            'conclusion': "L'établissement félicite chaleureusement l'élève et encourage les efforts de tous.",
        },
        'convocation_parent': {
            'intro': "Madame, Monsieur, vous êtes prié(e) de vous présenter au secrétariat de l'établissement :",
            'fields': [
                ('Élève concerné(e)', 'FALL Cheikh — 3ème B'),
                ('Objet', 'Comportement en classe'),
                ('Date de convocation', 'Jeudi 15 octobre 2026 à 10h00'),
                ('Personne à voir', 'Le Proviseur adjoint'),
                ('Lieu', 'Bureau de la Direction'),
            ],
            'conclusion': "Votre présence est indispensable au bon suivi de la scolarité de votre enfant.",
        },
        'recu_paiement': {
            'intro': "Reçu officiel de paiement des frais de scolarité :",
            'fields': [
                ('N° de reçu', 'REC-2026-0158'),
                ('Payeur', 'M. DIALLO (parent de Awa DIALLO, 6ème A)'),
                ('Montant réglé', '75 000 FCFA'),
                ('Mode de paiement', 'Orange Money'),
                ('Date', '05/10/2026'),
            ],
            'conclusion': "Ce reçu tient lieu de justificatif officiel de règlement. Merci de votre confiance.",
        },
        'bulletin': {
            'intro': "Extrait du bulletin de notes de la période :",
            'fields': [
                ('Élève', 'BARRY Aïssata — 5ème A'),
                ('Moyenne générale', '13,50 / 20'),
                ('Rang', '5 / 28'),
                ('Mention', 'Assez Bien'),
                ('Période', '1er Trimestre — 2025-2026'),
            ],
            'conclusion': "Bulletin complet disponible depuis le menu Bulletins de l'application.",
        },
        'demande_conge': {
            'intro': "Demande de congé adressée à la Direction de l'établissement :",
            'fields': [
                ('Demandeur', 'M. MAIGA Oumar — Enseignant'),
                ('Type de congé', 'Congé annuel'),
                ('Du', '20/12/2026'),
                ('Au', '04/01/2027'),
                ('Motif', 'Congés familiaux'),
            ],
            'conclusion': "Signé le demandeur, avis du supérieur hiérarchique et décision de la Direction.",
        },
    }

DOC_SAMPLE_DEFAULT = {
        'intro': "Exemplaire de démonstration du document.",
        'fields': [('Document', '__DOC_TITLE__')],
        'conclusion': "Ceci est un exemplaire fictif généré à titre d'illustration.",
    }


def index_documents(request):
    """Page Documents : catalogue professionnel avec exemplaire par type."""
    modeles = DocumentModele.objects.filter(est_actif=True)

    context = {
        'catalog': DOCUMENTS_CATALOG,
        'modeles_rh': modeles.filter(categorie='RH'),
        'modeles_admin': modeles.filter(categorie='ADMINISTRATIF'),
        'modeles_pedagogy': modeles.filter(categorie='PEDAGOGIQUE'),
    }
    return render(request, 'documents/documents_list.html', context)


@login_required
def document_sample_pdf(request, key):
    """Exemplaire PDF fictif pour chaque type de document du catalogue."""
    doc = next((d for d in DOCUMENTS_CATALOG if d['key'] == key), None)
    if doc is None:
        messages.error(request, "Type de document inconnu.")
        return redirect('documents:index')

    school = None
    from school_settings.models import SchoolSetting
    try:
        school = SchoolSetting.load()
    except Exception:
        school = None

    if key in DOC_SAMPLE_DATA:
        sample_data = DOC_SAMPLE_DATA[key]
    else:
        sample_data = dict(DOC_SAMPLE_DEFAULT)
        sample_data['fields'] = [('Document', doc['title'])]

    context = {
        'doc': doc,
        'school': school,
        'intro': sample_data['intro'],
        'fields': sample_data['fields'],
        'conclusion': sample_data['conclusion'],
        'reference': f"EX-2026/{doc['key'][:8].upper()}-0001",
        'date_str': timezone.localdate().strftime('%d/%m/%Y'),
    }

    html = render_to_string('documents/sample_pdf.html', context)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode('UTF-8')), result, encoding='UTF-8')
    if pdf.err:
        return HttpResponse("Erreur lors de la génération de l'exemplaire.", status=500)
    response = HttpResponse(result.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="Exemple_{key}.pdf"'
    return response


def generate_certificat_scolarite(request):
    """
    Génération du Certificat de Scolarité
    """
    if request.method == 'POST':
        form = CertificatScolariteForm(request.POST)
        if form.is_valid():
            context = {
                'eleve': {
                    'nom': form.cleaned_data['eleve_nom'],
                    'prenom': form.cleaned_data['eleve_prenom'],
                    'matricule': form.cleaned_data['matricule'],
                    'date_naissance': form.cleaned_data['date_naissance'].strftime('%d/%m/%Y'),
                    'classe': form.cleaned_data['classe'],
                },
                'annee_scolaire': form.cleaned_data['annee_scolaire'],
                'numero_reference': "2026/CS-0042",
            }
            # Rend directement le template imprimable/PDF
            return render(request, 'documents/certificat_scolarite_pdf.html', context)
    else:
        form = CertificatScolariteForm()

    return render(request, 'documents/form_certificat.html', {'form': form, 'titre': 'Certificat de Scolarité'})


def generate_certificat_frequentation(request):
    """Certificat de fréquentation : renvoie l'exemplaire PDF du catalogue."""
    return redirect('documents:document_sample', key='certificat_frequentation')


def generate_certificat_transfert(request):
    """Certificat de transfert : renvoie l'exemplaire PDF du catalogue."""
    return redirect('documents:document_sample', key='certificat_transfert')


def registre_presence(request):
    """Registre de présence : renvoie l'exemplaire PDF du catalogue."""
    return redirect('documents:document_sample', key='registre_presence')


@login_required
def demande_conge_create(request):
    """
    Soumettre une demande de congé/autorisation d'absence
    """
    if request.method == 'POST':
        form = DemandeCongeForm(request.POST)
        if form.is_valid():
            conge = form.save(commit=False)
            conge.employe = request.user
            conge.save()
            messages.success(request, "Votre demande de congé a été enregistrée avec succès.")
            return redirect('documents:index')
    else:
        form = DemandeCongeForm()

    return render(request, 'documents/demande_conge_form.html', {'form': form})


# ---------------------------------------------------------------------------
# Génération Word (.docx) : exemple professionnel au format PAYSAGE A4.
# ---------------------------------------------------------------------------
DOCX_NAVY = RGBColor(0x1D, 0x35, 0x57)
DOCX_RED = RGBColor(0xE6, 0x39, 0x46)
DOCX_BLUE = RGBColor(0x45, 0x7B, 0x9D)
DOCX_GRAY = RGBColor(0x64, 0x74, 0x8B)
DOCX_DARK = RGBColor(0x1E, 0x29, 0x3B)


def _docx_cell_shade(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def _docx_table_borders(table, color_hex='CBD5E1', sz='4'):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:%s' % edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:color'), color_hex)
        borders.append(el)
    tblPr.append(borders)


def _docx_para_rule(paragraph, color='1D3557', sz='18'):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _docx_run(p, text, size=10, bold=False, color=None, italic=False):
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


@login_required
def document_sample_word(request, key):
    """Exemple du document au format Word (.docx) — paysage A4 professionnel."""
    doc = next((d for d in DOCUMENTS_CATALOG if d['key'] == key), None)
    if doc is None:
        messages.error(request, "Type de document inconnu.")
        return redirect('documents:index')

    from school_settings.models import SchoolSetting
    try:
        school = SchoolSetting.load()
    except Exception:
        school = None
    if key in DOC_SAMPLE_DATA:
        data = DOC_SAMPLE_DATA[key]
    else:
        data = dict(DOC_SAMPLE_DEFAULT)
        data['fields'] = [('Document', doc['title'])]

    school_name = getattr(school, 'school_name', None) or "Gestion d'École"
    school_address = getattr(school, 'address', '') or ''
    school_phone = getattr(school, 'phone', '') or ''
    reference = "EX-2026/%s-0001" % doc['key'][:8].upper()
    date_str = timezone.localdate().strftime('%d/%m/%Y')
    ts_str = timezone.localtime().strftime('%d/%m/%Y %H:%M')

    document = DocxDocument()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.5)

    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)

    # ---- En-tête : école à gauche, titre du document à droite
    header = document.add_table(rows=1, cols=2)
    header.autofit = False
    left, right = header.rows[0].cells
    left.width = Cm(15.5)
    right.width = Cm(10.6)
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    _docx_run(p, school_name, size=16, bold=True, color=DOCX_NAVY)
    meta = school_address
    if school_phone:
        meta = ("%s · Tél : %s" % (meta, school_phone)) if meta else ("Tél : %s" % school_phone)
    p = left.add_paragraph()
    _docx_run(p, meta, size=9, color=DOCX_GRAY)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    _docx_run(p, doc['title'].upper(), size=14, bold=True, color=DOCX_RED)
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _docx_run(p, "Réf : %s" % reference, size=9.5, bold=True, color=DOCX_BLUE)

    rule = document.add_paragraph()
    _docx_para_rule(rule)
    rule.paragraph_format.space_after = Pt(10)

    # ---- Introduction
    intro = document.add_paragraph()
    _docx_run(intro, data['intro'], size=11, color=DOCX_DARK)
    intro.paragraph_format.space_after = Pt(10)

    # ---- Tableau d'identité (libellé / valeur)
    fields = data['fields']
    table = document.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _docx_table_borders(table)
    for i, (label, value) in enumerate(fields):
        c1, c2 = table.rows[i].cells
        c1.width = Cm(7.5)
        c2.width = Cm(18.6)
        p = c1.paragraphs[0]
        _docx_run(p, label, size=10, bold=True, color=DOCX_GRAY)
        _docx_cell_shade(c1, 'F8F9FA')
        p = c2.paragraphs[0]
        _docx_run(p, str(value), size=10, color=DOCX_DARK)
        for cell in (c1, c2):
            for par in cell.paragraphs:
                par.paragraph_format.space_before = Pt(3)
                par.paragraph_format.space_after = Pt(3)

    # ---- Conclusion
    concl = document.add_paragraph()
    concl.paragraph_format.space_before = Pt(12)
    _docx_run(concl, data['conclusion'], size=10.5, italic=True,
              color=RGBColor(0x33, 0x41, 0x55))

    # ---- Signatures : cachet à gauche, Directeur à droite
    sig = document.add_table(rows=1, cols=2)
    sig.autofit = False
    s1, s2 = sig.rows[0].cells
    s1.width = Cm(13)
    s2.width = Cm(13)
    p = s1.paragraphs[0]
    p.paragraph_format.space_before = Pt(24)
    _docx_run(p, "Cachet de l'établissement", size=10, color=RGBColor(0x33, 0x41, 0x55))
    p = s2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(24)
    _docx_run(p, "Le Directeur,", size=10, color=RGBColor(0x33, 0x41, 0x55))
    for _ in range(4):
        p = s2.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = s2.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _docx_run(p, "Nom et signature", size=8.5, color=DOCX_GRAY)

    # ---- Pied de page
    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(16)
    _docx_run(footer, "Document généré le %s — %s" % (ts_str, school_name),
              size=8, color=RGBColor(0x94, 0xA3, 0xB8))

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = 'attachment; filename="Exemple_%s.docx"' % key
    document.save(response)
    return response
