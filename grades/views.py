"""
Vues du module Notes & Bulletins :
- Liste des évaluations + import Excel
- Saisie des notes par classe / matière (grille)
- Bulletins (moyennes, coefficients, rangs) + exports PDF & Excel
"""
import os
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django_htmx.http import HttpResponseClientRedirect
from xhtml2pdf import pisa

from classes.models import Class
from courses.models import Subject
from school_settings.models import GeneralSetting, SchoolSetting
from students.models import Student

from .forms import PeriodForm, EvaluationForm, ExcelImportForm
from .models import AcademicYear, Evaluation, Grade, Period
from .services import (
    appreciation_for,
    class_ranking,
    general_average,
    mention_for,
    subject_breakdown,
)


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def _school():
    """Informations de l'établissement (singleton SchoolSetting)."""
    try:
        return SchoolSetting.load()
    except Exception:
        return None


def _bulletin_header_context(school_setting):
    """Contexte d'en-tête des bulletins : infos onglet École (avec repli),
    logo résolu sur disque et tutelles / république (Paramètres → Général)."""
    school = getattr(school_setting, 'school', None)
    logo_path = None
    if school_setting is not None and getattr(school_setting, 'logo', None):
        try:
            if school_setting.logo.storage.exists(school_setting.logo.name):
                logo_path = school_setting.logo.path
        except Exception:
            logo_path = None
    ctx = {
        'school_name': (getattr(school_setting, 'school_name', None)
                        or getattr(school, 'name', None) or "Gestion d'École"),
        'school_address': (getattr(school_setting, 'address', None)
                           or getattr(school, 'address', None)),
        'school_phone': (getattr(school_setting, 'phone', None)
                         or getattr(school, 'phone', None)),
        'school_email': (getattr(school_setting, 'email', None)
                         or getattr(school, 'email', None)),
        'logo_path': logo_path,
        'general': GeneralSetting.load(),
    }
    if school_setting is not None and getattr(school_setting, 'logo', None):
        try:
            ctx['logo_url'] = school_setting.logo.url
        except Exception:
            ctx['logo_url'] = None
    return ctx


def _render_pdf(template_src, context, filename):
    """Génère un PDF à partir d'un template HTML (xhtml2pdf)."""
    html = render_to_string(template_src, context)
    result = BytesIO()

    def _link(uri, rel):
        # Ressources locales (logo de l'école) résolues sur le disque
        return uri if uri and os.path.exists(uri) else None

    pdf = pisa.pisaDocument(
        BytesIO(html.encode('UTF-8')), result, encoding='UTF-8',
        link_callback=_link)
    if pdf.err:
        return HttpResponse(
            "Erreur lors de la génération du PDF.", status=500
        )
    response = HttpResponse(result.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# --------------------------------------------------------------------------
# Évaluations & liste des notes
# --------------------------------------------------------------------------

@login_required
def grade_list(request):
    """Vue principale : liste des évaluations et formulaire d'import Excel."""
    evaluations = Evaluation.objects.select_related(
        'classroom', 'subject', 'period', 'period__academic_year'
    ).all()

    classroom_id = request.GET.get('classroom')
    subject_id = request.GET.get('subject')
    period_id = request.GET.get('period')
    if classroom_id:
        evaluations = evaluations.filter(classroom_id=classroom_id)
    if subject_id:
        evaluations = evaluations.filter(subject_id=subject_id)
    if period_id:
        evaluations = evaluations.filter(period_id=period_id)

    context = {
        'evaluations': evaluations,
        'form': ExcelImportForm(),
        'classrooms': Class.objects.filter(is_active=True),
        'subjects': Subject.objects.all(),
        'periods': Period.objects.select_related('academic_year'),
        'f_classroom': classroom_id or '',
        'f_subject': subject_id or '',
        'f_period': period_id or '',
    }
    return render(request, 'grades/grade_list.html', context)


def _current_academic_year(school):
    """Année académique active de l'école, créée au besoin depuis les
    Paramètres → Général (année scolaire en cours)."""
    from datetime import date as _date
    from school_settings.models import GeneralSetting

    year = AcademicYear.objects.filter(school=school, is_active=True).first()
    if year:
        return year

    today = _date.today()
    default_name = GeneralSetting.load().current_academic_year
    if not default_name:
        default_name = f"{today.year}-{today.year + 1}" if today.month >= 9 \
            else f"{today.year - 1}-{today.year}"
    year, _created = AcademicYear.objects.get_or_create(
        school=school, name=default_name,
        defaults=dict(start_date=_date(today.year, 9, 1),
                      end_date=_date(today.year + 1, 6, 30), is_active=True),
    )
    # une seule année active à la fois
    AcademicYear.objects.filter(school=school).exclude(pk=year.pk) \
        .update(is_active=False)
    year.is_active = True
    year.save()
    return year


@login_required
def period_create(request):
    """Ajout d'une période (trimestre, semestre…) en modale."""
    school = getattr(request.user, 'school', None)
    year = _current_academic_year(school)
    initial = {}
    if year:
        initial['academic_year'] = year

    form = PeriodForm(request.POST or None, initial=initial)
    if request.method == 'POST' and form.is_valid():
        period = form.save(commit=False)
        period.school = school
        period.save()
        messages.success(request, f"Période « {period.name} » créée.")
        target = (request.POST.get('next') or request.GET.get('next')
                  or reverse('grades:grade_entry'))
        if request.htmx:
            return HttpResponseClientRedirect(target)
        return redirect(target)

    if request.htmx:
        return render(request, 'includes/form_modal.html', {
            'form': form,
            'modal_title': 'Nouvelle période',
            'modal_action': request.get_full_path(),
        })
    return render(request, 'grades/evaluation_form.html', {
        'form': form,
        'title': 'Nouvelle période',
        'next_url': '',
    })


@login_required
def evaluation_create(request):
    """Création d'une évaluation, puis retour à la saisie des notes."""
    next_url = request.GET.get('next') or request.POST.get('next')
    initial = {}
    for param in ('classroom', 'subject', 'period'):
        value = request.GET.get(param)
        if value:
            initial[param] = value
    default_date = request.GET.get('date') or None
    if default_date:
        initial['date'] = default_date

    form = EvaluationForm(request.POST or None, initial=initial)
    if request.method == 'POST' and form.is_valid():
        evaluation = form.save()
        messages.success(request, f"Évaluation « {evaluation.title} » créée avec succès.")
        target = next_url or reverse('grades:grade_entry_eval', args=[evaluation.pk])
        if request.htmx:
            return HttpResponseClientRedirect(target)
        return redirect(target)

    if request.htmx:
        return render(request, 'includes/form_modal.html', {
            'form': form,
            'modal_title': 'Nouvelle évaluation',
            'modal_action': request.get_full_path(),
        })
    return render(request, 'grades/evaluation_form.html', {
        'form': form,
        'title': 'Nouvelle évaluation',
        'next_url': next_url or '',
    })


@login_required
def evaluation_edit(request, pk):
    evaluation = get_object_or_404(Evaluation, pk=pk)
    form = EvaluationForm(request.POST or None, instance=evaluation)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, "Évaluation mise à jour.")
        if request.htmx:
            return HttpResponseClientRedirect(reverse('grades:grade_list'))
        return redirect('grades:grade_list')

    if request.htmx:
        return render(request, 'includes/form_modal.html', {
            'form': form,
            'modal_title': f"Modifier : {evaluation.title}",
            'modal_action': request.get_full_path(),
        })
    return render(request, 'grades/evaluation_form.html', {
        'form': form,
        'title': f"Modifier : {evaluation.title}",
        'next_url': '',
    })


@login_required
def evaluation_delete(request, pk):
    evaluation = get_object_or_404(Evaluation, pk=pk)
    if request.method == 'POST':
        title = evaluation.title
        evaluation.delete()
        messages.success(request, f"Évaluation « {title} » supprimée.")
        if request.htmx:
            return HttpResponseClientRedirect(reverse('grades:grade_list'))
        return redirect('grades:grade_list')

    if request.htmx:
        return render(request, 'includes/delete_modal.html', {
            'object': evaluation,
            'modal_title': 'Supprimer cette évaluation',
            'modal_action': request.get_full_path(),
        })
    return render(request, 'grades/evaluation_confirm_delete.html', {
        'evaluation': evaluation,
    })


# --------------------------------------------------------------------------
# Saisie des notes par classe / matière
# --------------------------------------------------------------------------

@login_required
def grade_entry(request, evaluation=None):
    """
    Saisie des notes :
    1. L'utilisateur choisit classe + matière + période → liste des évaluations.
    2. Il choisit (ou crée) une évaluation → grille de saisie des notes.
    """
    classrooms = Class.objects.filter(is_active=True).order_by('name')
    subjects = Subject.objects.all().order_by('name')
    periods = Period.objects.select_related('academic_year').order_by('-academic_year', 'id')

    evaluation_obj = None
    evaluations = None
    rows = []

    evaluation_id = evaluation or request.GET.get('evaluation') \
        or request.POST.get('evaluation')

    # --- Enregistrement de la grille -----------------------------------
    if request.method == 'POST' and evaluation_id:
        evaluation_obj = get_object_or_404(
            Evaluation.objects.select_related('classroom', 'subject'),
            pk=evaluation_id,
        )
        saved, errors = 0, 0
        with transaction.atomic():
            for student in evaluation_obj.classroom.students.filter(is_active=True):
                raw = (request.POST.get(f'score_{student.id}') or '').strip().replace(',', '.')
                appr = (request.POST.get(f'appreciation_{student.id}') or '').strip()[:255]

                if raw == '':
                    score = None
                else:
                    try:
                        score = float(raw)
                        if score < 0 or score > float(evaluation_obj.max_score):
                            raise ValueError
                    except ValueError:
                        errors += 1
                        messages.warning(
                            request,
                            f"Note invalide pour {student.full_name} "
                            f"(doit être entre 0 et {evaluation_obj.max_score})."
                        )
                        continue

                Grade.objects.update_or_create(
                    student=student,
                    evaluation=evaluation_obj,
                    defaults={'score': score, 'appreciation': appr},
                )
                saved += 1

        if saved:
            messages.success(
                request,
                f"{saved} note(s) enregistrée(s) pour « {evaluation_obj.title} »."
            )
        return redirect('grades:grade_entry_eval', evaluation=evaluation_obj.pk)

    # --- Affichage -------------------------------------------------------
    if evaluation_id:
        evaluation_obj = get_object_or_404(
            Evaluation.objects.select_related('classroom', 'subject', 'period',
                                              'period__academic_year'),
            pk=evaluation_id,
        )
        grades_by_student = {
            g.student_id: g for g in Grade.objects.filter(evaluation=evaluation_obj)
        }
        rows = [
            {
                'student': s,
                'grade': grades_by_student.get(s.id),
            }
            for s in evaluation_obj.classroom.students.filter(is_active=True)
        ]
    else:
        classroom_id = request.GET.get('classroom')
        subject_id = request.GET.get('subject')
        period_id = request.GET.get('period')
        filters = {}
        if classroom_id:
            filters['classroom_id'] = classroom_id
        if subject_id:
            filters['subject_id'] = subject_id
        if period_id:
            filters['period_id'] = period_id
        if filters:
            evaluations = Evaluation.objects.filter(**filters).select_related(
                'classroom', 'subject', 'period'
            ).order_by('-date')

    context = {
        'classrooms': classrooms,
        'subjects': subjects,
        'periods': periods,
        'evaluation': evaluation_obj,
        'evaluations': evaluations,
        'rows': rows,
        'f_classroom': request.GET.get('classroom', ''),
        'f_subject': request.GET.get('subject', ''),
        'f_period': request.GET.get('period', ''),
    }
    return render(request, 'grades/grade_entry.html', context)


@login_required
def import_grades_excel(request):
    """Traitement du fichier Excel téléversé (matricule, note, appréciation)."""
    if request.method == 'POST':
        form = ExcelImportForm(request.POST, request.FILES)
        if form.is_valid():
            excel_file = request.FILES['excel_file']
            evaluation = form.cleaned_data['evaluation']

            try:
                wb = openpyxl.load_workbook(excel_file, data_only=True)
                sheet = wb.active

                created_count = updated_count = 0
                errors = []

                with transaction.atomic():
                    for row_idx, row in enumerate(
                        sheet.iter_rows(min_row=2, values_only=True), start=2
                    ):
                        if not row or not row[0]:
                            continue

                        matricule = str(row[0]).strip()
                        raw_score = row[1] if len(row) > 1 else None
                        appreciation = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""

                        score = None
                        if raw_score is not None and str(raw_score).strip() != "":
                            try:
                                score = float(str(raw_score).replace(',', '.'))
                                if score < 0 or score > float(evaluation.max_score):
                                    errors.append(
                                        f"Ligne {row_idx}: note {score} hors barème "
                                        f"({evaluation.max_score})."
                                    )
                                    continue
                            except ValueError:
                                errors.append(
                                    f"Ligne {row_idx}: format de note invalide ('{raw_score}')."
                                )
                                continue

                        try:
                            student = Student.objects.get(student_id=matricule)
                        except Student.DoesNotExist:
                            errors.append(
                                f"Ligne {row_idx}: aucun élève avec le matricule '{matricule}'."
                            )
                            continue

                        grade, created = Grade.objects.update_or_create(
                            student=student,
                            evaluation=evaluation,
                            defaults={'score': score, 'appreciation': appreciation},
                        )
                        if created:
                            created_count += 1
                        else:
                            updated_count += 1

                if created_count or updated_count:
                    messages.success(
                        request,
                        f"Importation réussie : {created_count} créée(s), "
                        f"{updated_count} mise(s) à jour."
                    )
                else:
                    messages.warning(request, "Aucune note importée.")
                for error in errors[:5]:
                    messages.warning(request, error)

            except Exception as e:
                messages.error(request, f"Erreur lors de la lecture du fichier : {e}")

            return redirect('grades:grade_list')

    return redirect('grades:grade_list')


# --------------------------------------------------------------------------
# Bulletins
# --------------------------------------------------------------------------

@login_required
def bulletin_select(request):
    """Choix de l'élève (ou de la classe) et de la période → bulletin."""
    students = Student.objects.filter(is_active=True).select_related('class_group')
    periods = Period.objects.select_related('academic_year').all()
    classrooms = Class.objects.filter(is_active=True)

    student_id = request.GET.get('student')
    period_id = request.GET.get('period')
    if student_id and period_id:
        return redirect('grades:student_report_card',
                        student_id=student_id, period_id=period_id)

    class_id = request.GET.get('classroom')
    if class_id and period_id:
        return redirect('grades:export_class_bulletins_pdf',
                        classroom_id=class_id, period_id=period_id)

    return render(request, 'grades/bulletin_select.html', {
        'students': students,
        'periods': periods,
        'classrooms': classrooms,
    })


@login_required
def student_report_card(request, student_id, period_id):
    """Bulletin détaillé d'un élève : moyennes par matière, rang, mention."""
    student = get_object_or_404(
        Student.objects.select_related('class_group', 'class_group__program'),
        pk=student_id,
    )
    period = get_object_or_404(Period.objects.select_related('academic_year'), pk=period_id)

    rows = subject_breakdown(student, period)
    general_avg = general_average(rows)

    rank = None
    total_ranked = 0
    class_avg = None
    if student.class_group:
        results, ranks, class_avg = class_ranking(student.class_group, period)
        rank = ranks.get(student.id)
        total_ranked = sum(1 for r in results if r['rank'] is not None)

    total_coef = sum(r['coefficient'] for r in rows)
    total_points = sum(r['points'] for r in rows if r['points'] is not None)

    school_setting = _school()
    context = {
        'student': student,
        'period': period,
        'school': school_setting,
        'rows': rows,
        'general_average': general_avg,
        'mention': mention_for(general_avg),
        'rank': rank,
        'total_ranked': total_ranked,
        'class_average': class_avg,
        'total_coef': total_coef,
        'total_points': round(total_points, 2),
    }
    context.update(_bulletin_header_context(school_setting))
    return render(request, 'grades/student_report_card.html', context)


@login_required
def export_bulletin_excel(request, student_id, period_id):
    """Export Excel du bulletin d'un élève."""
    student = get_object_or_404(
        Student.objects.select_related('class_group', 'class_group__program'),
        pk=student_id,
    )
    period = get_object_or_404(Period.objects.select_related('academic_year'), pk=period_id)

    rows = subject_breakdown(student, period)
    general_avg = general_average(rows)
    total_coef = sum(r['coefficient'] for r in rows)

    rank = class_avg = None
    total_ranked = 0
    if student.class_group:
        results, ranks, class_avg = class_ranking(student.class_group, period)
        rank = ranks.get(student.id)
        total_ranked = sum(1 for r in results if r['rank'] is not None)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bulletin"

    school_setting = _school()
    school = school_setting
    ctx = _bulletin_header_context(school_setting)

    font_name = Font(name='Arial', size=16, bold=True, color='065F46')
    font_rep = Font(name='Arial', size=10, bold=True, color='0F172A')
    font_band = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    font_sub = Font(name='Arial', size=9, color='FFFFFF')
    font_th = Font(name='Arial', size=9, bold=True, color='FFFFFF')
    font_lbl = Font(name='Arial', size=8, bold=True, color='64748B')
    font_big = Font(name='Arial', size=13, bold=True, color='1D3557')
    fill_band = PatternFill(start_color='1D3557', end_color='1D3557', fill_type='solid')
    fill_th = PatternFill(start_color='1D3557', end_color='1D3557', fill_type='solid')
    fill_total = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    fill_soft = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'),
    )
    centered = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_al = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # --- En-tête officiel : tutelles | école | république + devise ---
    general = ctx['general']
    ws.merge_cells('A1:B1')
    c = ws['A1']
    c.value = (general.ministry or '')
    c.font = Font(name='Arial', size=8, bold=True, color='0F172A')
    c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    if general.directorate:
        ws.merge_cells('A2:B2')
        ws['A2'] = general.directorate
        ws['A2'].font = Font(name='Arial', size=8, color='0F172A')
        ws['A2'].alignment = Alignment(horizontal='left', vertical='top')

    ws.merge_cells('C1:C2')
    c = ws['C1']
    c.value = ctx['school_name']
    c.font = font_name
    c.alignment = centered

    ws.merge_cells('D1:E1')
    c = ws['D1']
    c.value = (general.country_name or '')
    c.font = font_rep
    c.alignment = centered
    ws.merge_cells('D2:E2')
    c = ws['D2']
    c.value = (general.motto or '')
    c.font = Font(name='Arial', size=9, color='0F172A')
    c.alignment = centered

    # --- Bandeau titre ---
    ws.merge_cells('A4:E4')
    c = ws['A4']
    c.value = f"BULLETIN DE NOTES — {period.name} — Année scolaire {period.academic_year.name}"
    c.font = font_band
    c.fill = fill_band
    c.alignment = centered
    ws.row_dimensions[4].height = 24

    # --- Identité élève ---
    ws.merge_cells('A5:B5')
    ws['A5'] = f"Élève : {student.last_name} {student.first_name}"
    ws['A5'].font = Font(name='Arial', size=9, bold=True)
    ws.merge_cells('C5:D5')
    ws['C5'] = f"Matricule : {student.student_id}"
    ws['C5'].font = Font(name='Arial', size=9)
    ws['E5'] = f"Classe : {student.class_group.name if student.class_group else 'N/A'}"
    ws['E5'].font = Font(name='Arial', size=9)

    # --- Tableau des notes ---
    headers = ["Matière", "Coef.", "Moyenne / 20", "Points", "Appréciation"]
    ws.append([])
    ws.append(headers)
    for col_num in range(1, 6):
        cell = ws.cell(row=7, column=col_num)
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = centered
        cell.border = thin_border
    ws.row_dimensions[7].height = 18

    current_row = 8
    for row in rows:
        ws.append([
            row['subject'].name,
            row['coefficient'],
            row['average'] if row['average'] is not None else "—",
            row['points'] if row['points'] is not None else "—",
            row['appreciation'] or "—",
        ])
        for col_num in range(1, 6):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = thin_border
            cell.alignment = centered if col_num > 1 else left_al
        current_row += 1

    total_points = round(sum(r['points'] for r in rows if r['points']), 2)
    ws.append(["TOTAL", total_coef, "—", total_points, "—"])
    for col_num in range(1, 6):
        cell = ws.cell(row=current_row, column=col_num)
        cell.font = Font(name='Arial', size=9, bold=True)
        cell.fill = fill_total
        cell.border = thin_border
        cell.alignment = centered if col_num > 1 else left_al
    current_row += 1

    ws.append([
        "MOYENNE GÉNÉRALE", "",
        (general_avg if general_avg is not None else "N/A"),
        mention_for(general_avg), "",
    ])
    ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
    for col_num in range(1, 6):
        cell = ws.cell(row=current_row, column=col_num)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = fill_band
        cell.border = thin_border
        cell.alignment = centered
    ws.row_dimensions[current_row].height = 20
    current_row += 2

    # --- Synthèse : rang / moyenne classe / mention ---
    ws.append([])
    summary_row = current_row
    rank_txt = f"{rank}/{total_ranked}" if rank else "—"
    blocks = [
        (1, 2, 'RANG', rank_txt),
        (3, 3, 'MOYENNE DE LA CLASSE', class_avg if class_avg is not None else '—'),
        (4, 5, 'MENTION', mention_for(general_avg)),
    ]
    for col_start, col_end, label, value in blocks:
        if col_end > col_start:
            ws.merge_cells(start_row=summary_row, start_column=col_start,
                           end_row=summary_row, end_column=col_end)
            ws.merge_cells(start_row=summary_row + 1, start_column=col_start,
                           end_row=summary_row + 1, end_column=col_end)
        lc = ws.cell(row=summary_row, column=col_start)
        lc.value = label
        lc.font = font_lbl
        lc.alignment = centered
        vc = ws.cell(row=summary_row + 1, column=col_start)
        vc.value = value
        vc.font = font_big
        vc.alignment = centered
        for r in (summary_row, summary_row + 1):
            for cc in range(col_start, col_end + 1):
                ws.cell(row=r, column=cc).border = thin_border
                ws.cell(row=r, column=cc).fill = fill_soft
    ws.row_dimensions[summary_row + 1].height = 22
    current_row = summary_row + 3

    # --- Signatures ---
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    ws.cell(row=current_row, column=1).value = "Signature du Chef d'établissement :"
    ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
    ws.cell(row=current_row, column=4).value = "Signature du Parent / Tuteur :"
    ws.cell(row=current_row, column=1).font = Font(name='Arial', size=9)
    ws.cell(row=current_row, column=4).font = Font(name='Arial', size=9)
    for r in (current_row + 3,):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        for col in (1, 4):
            cell = ws.cell(row=r, column=col)
            cell.border = Border(top=Side(style='thin', color='64748B'))
    current_row = r + 2

    # --- Pied de page infos école ---
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    foot_parts = [p for p in (ctx['school_address'],
                              f"Email : {ctx['school_email']}" if ctx['school_email'] else None,
                              f"Tél : {ctx['school_phone']}" if ctx['school_phone'] else None) if p]
    cell = ws.cell(row=current_row, column=1)
    cell.value = ' · '.join(foot_parts)
    cell.font = Font(name='Arial', size=8, color='64748B')
    cell.alignment = centered

    # --- Logo de l'école (même visuel que le PDF) ---
    if ctx.get('logo_path') and os.path.exists(ctx['logo_path']):
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(ctx['logo_path'])
            img.width, img.height = 40, 40
            ws.add_image(img, 'C1')
            ws.row_dimensions[1].height = 20
            ws.row_dimensions[2].height = 20
        except Exception:
            pass  # image optionnelle (Pillow absent, format illisible…)

    # --- Largeurs + MISE EN PAGE A5 (impression directe) ---
    widths = {'A': 26, 'B': 8, 'C': 13, 'D': 12, 'E': 26}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.page_setup.paperSize = 11  # A5
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.35, right=0.35, top=0.45, bottom=0.45)
    ws.print_area = f'A1:E{current_row}'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    safe_name = student.last_name or 'eleve'
    response['Content-Disposition'] = (
        f'attachment; filename="Bulletin_{safe_name}_{period.name}.xlsx"'
    )
    wb.save(response)
    return response


@login_required
def export_bulletin_pdf(request, student_id, period_id):
    """Export PDF du bulletin d'un élève (xhtml2pdf)."""
    student = get_object_or_404(
        Student.objects.select_related('class_group', 'class_group__program'),
        pk=student_id,
    )
    period = get_object_or_404(Period.objects.select_related('academic_year'), pk=period_id)

    rows = subject_breakdown(student, period)
    general_avg = general_average(rows)

    rank = total_ranked = class_avg = None
    if student.class_group:
        results, ranks, class_avg = class_ranking(student.class_group, period)
        rank = ranks.get(student.id)
        total_ranked = sum(1 for r in results if r['rank'] is not None)

    school_setting = _school()
    context = {
        'school': school_setting,
        'student': student,
        'period': period,
        'rows': rows,
        'general_average': general_avg,
        'mention': mention_for(general_avg),
        'rank': rank,
        'total_ranked': total_ranked,
        'class_average': class_avg,
        'total_coef': sum(r['coefficient'] for r in rows),
        'total_points': round(sum(r['points'] for r in rows if r['points']), 2),
    }
    context.update(_bulletin_header_context(school_setting))
    filename = f"Bulletin_{student.student_id}_{period.name}.pdf"
    return _render_pdf('grades/report_card_pdf.html', context, filename)


@login_required
def export_class_bulletins_pdf(request, classroom_id, period_id):
    """Export PDF de tous les bulletins d'une classe (un par page)."""
    classroom = get_object_or_404(
        Class.objects.select_related('program'), pk=classroom_id
    )
    period = get_object_or_404(Period.objects.select_related('academic_year'), pk=period_id)

    results, ranks, class_avg = class_ranking(classroom, period)

    bulletins = []
    for item in results:
        student = item['student']
        rows = subject_breakdown(student, period)
        bulletins.append({
            'student': student,
            'rows': rows,
            'general_average': item['average'],
            'rank': item['rank'],
            'mention': mention_for(item['average']),
            'total_coef': sum(r['coefficient'] for r in rows),
            'total_points': round(sum(r['points'] for r in rows if r['points']), 2),
        })

    school_setting = _school()
    context = {
        'school': school_setting,
        'classroom': classroom,
        'period': period,
        'class_average': class_avg,
        'bulletins': bulletins,
    }
    context.update(_bulletin_header_context(school_setting))
    filename = f"Bulletins_{classroom.name}_{period.name}.pdf"
    return _render_pdf('grades/class_bulletins_pdf.html', context, filename)
